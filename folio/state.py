import asyncio
import io
import json
import math
from pathlib import Path
import re
from typing import Any, TypedDict
import urllib.parse
import uuid

from bs4 import BeautifulSoup
from docx import Document as WordDocument
from ebooklib import ITEM_DOCUMENT, epub
import httpx
import markdown
import reflex as rx

LANGUAGES = [
    {"code": "ml", "name": "Malayalam"},
    {"code": "en", "name": "English"},
    {"code": "hi", "name": "Hindi"},
    {"code": "ta", "name": "Tamil"},
    {"code": "te", "name": "Telugu"},
    {"code": "kn", "name": "Kannada"},
    {"code": "ar", "name": "Arabic"},
    {"code": "es", "name": "Spanish"},
    {"code": "fr", "name": "French"},
    {"code": "de", "name": "German"},
    {"code": "it", "name": "Italian"},
    {"code": "pt", "name": "Portuguese"},
    {"code": "ru", "name": "Russian"},
    {"code": "zh-CN", "name": "Chinese"},
    {"code": "ja", "name": "Japanese"},
]

# Backend storage for full page text to keep WebSocket payload tiny
_DOCUMENT_STORE: dict[str, list[list[str]]] = {}

class DocumentMeta(TypedDict):
    id: str
    title: str
    filename: str
    format: str
    words: int
    pages: int
    minutes: int

_HTTP_CLIENT = httpx.AsyncClient(
    timeout=15.0,
    headers={
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124.0.0.0 Safari/537.36"
    },
    follow_redirects=True,
    limits=httpx.Limits(max_keepalive_connections=20, max_connections=40),
)

async def _fast_google_translate(text: str, target: str) -> str:
    clean = text.strip()
    if not clean:
        return ""

    url = "https://translate.googleapis.com/translate_a/single"
    params = {
        "client": "gtx",
        "sl": "auto",
        "tl": target,
        "dt": "t",
        "q": clean,
    }

    try:
        res = await _HTTP_CLIENT.get(url, params=params)
        if res.status_code == 200:
            data = res.json()
            if data and isinstance(data, list) and data[0]:
                out = "".join(part[0] for part in data[0] if part and part[0])
                if out.strip():
                    return out.strip()
    except Exception as e:
        print(f"[FastTranslate Error]: {e}")

    try:
        url_dict = f"https://clients5.google.com/translate_a/t?client=dict-chrome-ex&sl=auto&tl={target}&q={urllib.parse.quote(clean)}"
        res = await _HTTP_CLIENT.get(url_dict)
        if res.status_code == 200:
            data = res.json()
            if isinstance(data, list) and len(data) > 0:
                return data[0] if isinstance(data[0], str) else data[0][0]
    except Exception:
        pass

    return clean

def _normalize(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")
    text = re.sub(r"[\t\u00a0 ]+", " ", text)
    return re.sub(r"\n{3,}", "\n\n", text).strip()

def _decode_text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-16", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")

def _html_to_text(source: str) -> str:
    soup = BeautifulSoup(source, "html.parser")
    for el in soup(["script", "style", "noscript", "svg", "nav", "footer", "header", "form"]):
        el.decompose()
    for el in soup.find_all(["p", "div", "section", "article", "h1", "h2", "h3", "h4", "h5", "h6", "li", "blockquote", "br"]):
        el.insert_before("\n\n")
    return _normalize(soup.get_text(separator=""))

def _paragraphs(text: str) -> list[str]:
    return [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]

def _paginate(text: str) -> list[list[str]]:
    pages, current, count = [], [], 0
    for para in _paragraphs(text):
        words = para.split()
        while words:
            space = 250 - count
            if count and (len(words) > space or (count >= 180 and len(words) > 70)):
                pages.append(current)
                current, count, space = [], 0, 250
            take = min(len(words), space)
            current.append(" ".join(words[:take]))
            words = words[take:]
            count += take
            if words:
                pages.append(current)
                current, count = [], 0
    if current:
        pages.append(current)
    return pages or [[""]]

def _extract_pages_fast(data: bytes, suffix: str) -> list[list[str]]:
    """Extracts text rapidly without blocking UI loops."""
    if suffix == ".pdf":
        # try:
        #     import fitz  # PyMuPDF: C-accelerated
        #     doc = fitz.open(stream=data, filetype="pdf")
        #     pages = []
        #     for page in doc:
        #         t = page.get_text().strip()
        #         if t:
        #             paras = [p.strip() for p in re.split(r"\n\s*\n", t) if p.strip()]
        #             if paras:
        #                 pages.append(paras)
        #     doc.close()
        #     if pages:
        #         return pages
        # except Exception:
        #     pass
        try:
            import pymupdf  # Changed from fitz
            doc = pymupdf.open(stream=data, filetype="pdf")  # Changed from fitz.open
            pages = []
            for page in doc:
                t = page.get_text().strip()
                if t:
                    paras = [p.strip() for p in re.split(r"\n\s*\n", t) if p.strip()]
                    if paras:
                        pages.append(paras)
            doc.close()
            if pages:
                return pages
        except Exception:
            pass

        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        pages = []
        for p in reader.pages:
            t = _normalize(p.extract_text() or "")
            if t:
                paras = _paragraphs(t)
                if paras:
                    pages.append(paras)
        if not pages:
            raise ValueError("No readable text found in PDF.")
        return pages

    elif suffix == ".docx":
        doc = WordDocument(io.BytesIO(data))
        text = _normalize("\n\n".join(p.text for p in doc.paragraphs if p.text))
    elif suffix == ".epub":
        book = epub.read_epub(io.BytesIO(data))
        parts = [_html_to_text(item.get_content().decode("utf-8", errors="replace"))
                 for item in book.get_items_of_type(ITEM_DOCUMENT)]
        text = _normalize("\n\n".join(parts))
    elif suffix in (".html", ".htm"):
        text = _html_to_text(_decode_text(data))
    elif suffix in (".md", ".markdown"):
        text = _html_to_text(markdown.markdown(_decode_text(data)))
    else:
        text = _normalize(_decode_text(data))

    if not text:
        raise ValueError("File contains no readable text.")
    return _paginate(text)

def _build_docx(title: str, text: str) -> bytes:
    doc = WordDocument()
    doc.add_heading(title, level=0)
    for p in text.split("\n\n"):
        if p.strip():
            doc.add_paragraph(p.strip())
    bio = io.BytesIO()
    doc.save(bio)
    bio.seek(0)
    return bio.getvalue()

class ReaderState(rx.State):
    documents: list[DocumentMeta] = []
    current_page_paragraphs: list[str] = []
    translations_cache: dict[str, str] = {}

    active_id: str = ""
    error_message: str = ""
    is_loading: bool = False
    font_size: int = 18
    line_spacing: str = "Relaxed"
    reading_width: str = "Comfortable"
    languages: list[dict[str, str]] = LANGUAGES
    target_language: str = "ml"
    selected_passage: int = -1
    is_dark: bool = False
    show_reader_window: bool = False
    mobile_shelf_open: bool = False
    mobile_view_tab: str = "split"
    translation_loading: bool = False
    translation_status: str = ""
    editing_id: str = ""
    editor_title: str = ""
    editor_text: str = ""
    current_page: int = 0

    export_format_reader: str = "docx"
    export_format_editor: str = "docx"

    def load_library(self):
        if self.documents and not self.active_id:
            self.select_document(self.documents[0]["id"])

    @rx.var
    def active_document(self) -> DocumentMeta:
        default_meta: DocumentMeta = {
            "id": "",
            "title": "",
            "filename": "",
            "format": "",
            "words": 0,
            "pages": 0,
            "minutes": 0,
        }
        for doc in self.documents:
            if doc.get("id") == self.active_id:
                return doc
        return default_meta

    @rx.var
    def visible_paragraphs(self) -> list[str]:
        return self.current_page_paragraphs

    @rx.var
    def page_start(self) -> int:
        pages = _DOCUMENT_STORE.get(self.active_id, [])
        return sum(len(page) for page in pages[:self.current_page])

    @rx.var
    def visible_translations(self) -> list[str]:
        paras = self.visible_paragraphs
        start = self.page_start
        doc_id = self.active_id
        res = []
        for idx in range(len(paras)):
            key = f"{doc_id}:{start + idx}"
            res.append(self.translations_cache.get(key, ""))
        return res

    @rx.var
    def current_page_translated_text(self) -> str:
        paras = self.visible_paragraphs
        start = self.page_start
        doc_id = self.active_id
        if not paras:
            return ""
        out = []
        for idx, p in enumerate(paras):
            key = f"{doc_id}:{start + idx}"
            trans = self.translations_cache.get(key)
            out.append(trans.strip() if trans and trans.strip() else f"[{p}]")
        return "\n\n".join(out)

    def _sync_current_page(self):
        pages = _DOCUMENT_STORE.get(self.active_id, [])
        if 0 <= self.current_page < len(pages):
            self.current_page_paragraphs = pages[self.current_page]
        else:
            self.current_page_paragraphs = []

    def toggle_theme(self):
        self.is_dark = not self.is_dark

    def toggle_reader_window(self):
        self.show_reader_window = not self.show_reader_window

    def toggle_mobile_shelf(self):
        self.mobile_shelf_open = not self.mobile_shelf_open

    def close_mobile_shelf(self):
        self.mobile_shelf_open = False

    def set_mobile_tab(self, tab: str):
        self.mobile_view_tab = tab

    def set_export_format_reader(self, fmt: str):
        self.export_format_reader = fmt

    def set_export_format_editor(self, fmt: str):
        self.export_format_editor = fmt

    def speak_text(self, text: str):
        if not text.strip():
            return
        safe_text = json.dumps(text.strip())
        target_lang = self.target_language
        lang_tag = f"{target_lang}-IN" if target_lang in ["ml", "hi", "ta", "te", "kn"] else target_lang
        script = f"""
        if ('speechSynthesis' in window) {{
            window.speechSynthesis.cancel();
            const utterance = new SpeechSynthesisUtterance({safe_text});
            utterance.lang = '{lang_tag}';
            utterance.rate = 0.95;
            window.speechSynthesis.speak(utterance);
        }}
        """
        return rx.call_script(script)

    def export_translation_from_reader(self):
        doc = self.active_document
        text = self.current_page_translated_text
        if not text:
            return
        title = f"{doc.get('title', 'Document')}_P{self.current_page + 1}_{self.target_language.upper()}"
        clean_name = title.replace(" ", "_")

        fmt = self.export_format_reader
        if fmt == "docx":
            file_bytes = _build_docx(title, text)
            return rx.download(data=file_bytes, filename=f"{clean_name}.docx")
        elif fmt == "md":
            md_content = f"# {title}\n\n{text}"
            return rx.download(data=md_content, filename=f"{clean_name}.md")
        elif fmt == "pdf":
            return rx.call_script("window.print();")
        else:
            return rx.download(data=text, filename=f"{clean_name}.txt")

    def export_editor_text(self):
        title = self.editor_title.strip() or "Untitled Document"
        content = self.editor_text.strip()
        clean_name = title.replace(" ", "_")

        fmt = self.export_format_editor
        if fmt == "docx":
            file_bytes = _build_docx(title, content)
            return rx.download(data=file_bytes, filename=f"{clean_name}.docx")
        elif fmt == "md":
            md_content = f"# {title}\n\n{content}"
            return rx.download(data=md_content, filename=f"{clean_name}.md")
        elif fmt == "pdf":
            return rx.call_script("window.print();")
        else:
            return rx.download(data=content, filename=f"{clean_name}.txt")

    def create_new_document(self):
        new_id = str(uuid.uuid4())
        doc: DocumentMeta = {
            "id": new_id,
            "title": "Untitled Document",
            "filename": "untitled.txt",
            "format": "TXT",
            "words": 0,
            "pages": 1,
            "minutes": 1,
        }
        _DOCUMENT_STORE[new_id] = [[""]]
        self.documents.append(doc)
        self.active_id = new_id
        self.current_page = 0
        self._sync_current_page()
        self.editing_id = new_id
        self.editor_title = "Untitled Document"
        self.editor_text = ""
        self.mobile_shelf_open = False
        return rx.redirect("/edit")

    def change_page(self, direction: int):
        nxt = self.current_page + direction
        total_pages = self.active_document.get("pages", 1)
        if 0 <= nxt < total_pages:
            self.current_page = nxt
            self._sync_current_page()
            self.selected_passage = -1

    def select_document(self, doc_id: str):
        self.active_id = doc_id
        self.current_page = 0
        self._sync_current_page()
        self.selected_passage = -1
        self.translation_status = ""
        self.mobile_shelf_open = False

    def remove_document(self, doc_id: str):
        self.documents = [d for d in self.documents if d.get("id") != doc_id]
        _DOCUMENT_STORE.pop(doc_id, None)
        if self.active_id == doc_id:
            if self.documents:
                self.select_document(self.documents[-1]["id"])
            else:
                self.active_id = ""
                self.current_page = 0
                self.current_page_paragraphs = []

    async def handle_upload(self, files: list[rx.UploadFile]):
        if not files:
            self.error_message = "No file selected."
            return
        self.is_loading = True
        self.error_message = ""
        yield

        new_docs = list(self.documents)
        last_id = ""

        for file in files:
            name = Path(file.filename or "file").name
            suffix = Path(name).suffix.lower()
            try:
                data = await file.read()
                # Run accelerated extraction in background thread
                pages = await asyncio.to_thread(_extract_pages_fast, data, suffix)
                total_words = sum(len(p.split()) for page in pages for p in page)

                doc_id = str(uuid.uuid4())
                _DOCUMENT_STORE[doc_id] = pages

                doc: DocumentMeta = {
                    "id": doc_id,
                    "title": Path(name).stem.replace("_", " ").title(),
                    "filename": name,
                    "format": suffix.lstrip(".").upper(),
                    "words": total_words,
                    "pages": len(pages),
                    "minutes": max(1, math.ceil(total_words / 220)),
                }
                new_docs.append(doc)
                last_id = doc_id
            except Exception as e:
                self.error_message = f"{name}: {str(e)}"
                print(f"[Upload Error]: {e}")

        self.documents = new_docs
        self.is_loading = False
        if last_id:
            self.select_document(last_id)
            self.mobile_shelf_open = False

        yield rx.clear_selected_files("documents")

    def open_editor(self):
        doc = self.active_document
        pages = _DOCUMENT_STORE.get(doc.get("id", ""), [])
        all_text = "\n\n".join(p for page in pages for p in page)
        self.editing_id = doc.get("id", "")
        self.editor_title = doc.get("title", "")
        self.editor_text = all_text
        return rx.redirect("/edit")

    def save_edit(self, form_data: dict[str, Any]):
        title = form_data.get("title", "").strip() or "Untitled Document"
        text = _normalize(form_data.get("text", "") or "")
        pages = _paginate(text) if text else [[""]]
        words = len(text.split())

        _DOCUMENT_STORE[self.editing_id] = pages

        for i, doc in enumerate(self.documents):
            if doc.get("id") == self.editing_id:
                self.documents[i].update({
                    "title": title,
                    "words": words,
                    "pages": len(pages),
                    "minutes": max(1, math.ceil(words / 220)),
                })
                break
        self._sync_current_page()
        return rx.redirect("/")

    def select_passage(self, index: int):
        self.selected_passage = index
        self.translation_status = f"Selected Passage {index + 1}."

    def set_target_language(self, code: str):
        self.target_language = code

    async def translate_selected_passage(self):
        paras = self.visible_paragraphs
        local_idx = self.selected_passage - self.page_start
        if local_idx < 0 or local_idx >= len(paras):
            self.translation_status = "⚠️ Click a passage on the left to select it."
            return

        self.translation_loading = True
        self.translation_status = f"Translating Passage {self.selected_passage + 1}..."
        yield

        target = self.target_language
        text = paras[local_idx]
        translated = await _fast_google_translate(text, target)

        self.translations_cache[f"{self.active_id}:{self.selected_passage}"] = translated
        self.translation_loading = False
        self.translation_status = "✅ Done."

    async def translate_entire_page(self):
        paras = self.visible_paragraphs
        if not paras:
            return

        self.translation_loading = True
        self.translation_status = "Translating current page..."
        yield

        target = self.target_language
        start_idx = self.page_start
        end_idx = start_idx + len(paras)
        indices = list(range(start_idx, end_idx))

        delimiter = "\n\n§§§\n\n"
        combined = delimiter.join(paras)

        result = await _fast_google_translate(combined, target)
        parts = [p.strip() for p in result.split("§§§")]

        if len(parts) == len(paras):
            for idx, text in zip(indices, parts):
                self.translations_cache[f"{self.active_id}:{idx}"] = text
        else:
            tasks = [_fast_google_translate(p, target) for p in paras]
            results = await asyncio.gather(*tasks)
            for idx, text in zip(indices, results):
                self.translations_cache[f"{self.active_id}:{idx}"] = text

        self.translation_loading = False
        self.translation_status = "✅ Page translated."

    def decrease_font(self): self.font_size = max(14, self.font_size - 1)
    def increase_font(self): self.font_size = min(28, self.font_size + 1)
    def toggle_width(self):
        self.reading_width = "Wide" if self.reading_width == "Comfortable" else "Comfortable"
    def cycle_spacing(self):
        self.line_spacing = {"Compact": "Relaxed", "Relaxed": "Airy", "Airy": "Compact"}[self.line_spacing]